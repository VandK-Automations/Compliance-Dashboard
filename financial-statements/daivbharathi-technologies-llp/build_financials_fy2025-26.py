# -*- coding: utf-8 -*-
"""Build Daivbharathi Technologies LLP FY 2025-26 financial statements in the Checkaro reference format."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter

OUT = 'Daivbharathi_Technologies_LLP_Financials_FY_2025-26.xlsx'
FONT = 'Segoe UI Semilight'
NF_AMT  = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"_);_(@_)'
NF_AMT2 = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"_);_(@_)'
NF_TB   = '_ * #,##0_ ;_ * \\-#,##0_ ;_ * "-"??_ ;_ @_ '
NF_TB2  = '_ * #,##0.00_ ;_ * \\-#,##0.00_ ;_ * "-"??_ ;_ @_ '
THIN = Side(style='thin'); MED = Side(style='medium'); NONE = Side(style=None)
YELLOW = PatternFill('solid', fgColor='FFFFFF00')
BLUEHDR = PatternFill('solid', fgColor='FF00B0F0')
GREY = PatternFill('solid', fgColor='FFF2F2F2')

def put(ws, ref, value=None, bold=False, italic=False, size=11, nf=None, h=None, v=None, wrap=None,
        indent=0, top=None, bottom=None, left=None, right=None, fill=None, color=None, name=FONT):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name=name, size=size, bold=bold, italic=italic, color=color)
    if nf: c.number_format = nf
    if h or v or wrap or indent:
        c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)
    if top or bottom or left or right:
        c.border = Border(top=top or NONE, bottom=bottom or NONE, left=left or NONE, right=right or NONE)
    if fill: c.fill = fill
    return c

def row_border(ws, row, cols, top=None, bottom=None):
    for col in cols:
        c = ws[f'{col}{row}']
        b = c.border
        c.border = Border(top=top or b.top, bottom=bottom or b.bottom, left=b.left, right=b.right)

def box(ws, ref_range):
    """thin box around each cell in range"""
    for row in ws[ref_range]:
        for c in row:
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

def setup_print(ws, area, landscape=False, fit_h=0):
    ws.print_area = area
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_h
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = 0.3937007874015748
    ws.sheet_view.showGridLines = False

def widths(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v

def heights(ws, d):
    for k, v in d.items():
        ws.row_dimensions[k].height = v

wb = Workbook()

# =====================================================================================
# MASTER
# =====================================================================================
ms = wb.active; ms.title = 'Master'
ms.sheet_view.showGridLines = False
widths(ms, {'A': 36.6, 'B': 44.0, 'C': 18.0, 'D': 18.0, 'E': 22.0})
TBP = '[To be provided]'
master_rows = [
    (2,  'Name of the LLP', 'Daivbharathi Technologies LLP', False),
    (3,  'Address', 'No 63, 1st Floor, 1st Main, 7th Block, Koramangala VI Block, Bangalore 560095', False),
    (4,  'LLPIN', TBP, True),
    (5,  'Period', '2025-26', False),
    (6,  'PAN', TBP, True),
    (7,  'Financial Year Opening Date', dt.datetime(2025, 4, 1), False),
    (8,  'Financial Year Closing Date', dt.datetime(2026, 3, 31), False),
    (9,  'Previous Year Closing Date', dt.datetime(2025, 3, 31), False),
    (10, 'Trial Balance period as per Tally (CY)', '1-Apr-25 to 23-Mar-26', False),
    (11, 'Trial Balance period as per Tally (PY)', '1-Apr-24 to 31-Mar-25', False),
    (13, "Auditor's Name", TBP, True),
    (14, 'Designation', TBP, True),
    (15, 'Membership No.', TBP, True),
    (16, 'Audit Firm Name', TBP, True),
    (17, 'Firm Registration No. (FRN)', TBP, True),
    (18, 'UDIN', TBP, True),
    (19, 'Date (Auditor)', TBP, True),
    (20, 'Place (Auditor)', TBP, True),
    (22, 'Designated Partner 1', 'Samartha Raghava Nagabhushanam', False),
    (23, 'DPIN No.', TBP, True),
    (24, 'Date', TBP, True),
    (25, 'Place', TBP, True),
    (26, 'Designated Partner 2', TBP, True),
    (27, 'DPIN No.', TBP, True),
    (28, 'Date', '=B24', False),
    (29, 'Place', TBP, True),
    (31, 'Financials Rounded Off by', 100, False),
]
for r, label, val, fillme in master_rows:
    put(ms, f'A{r}', label, bold=True, h='right', v='top' if r == 3 else None)
    c = put(ms, f'B{r}', val, h='left', wrap=(r == 3), fill=YELLOW if fillme else None)
    if isinstance(val, dt.datetime):
        c.number_format = 'dd-mm-yyyy'
ms.row_dimensions[3].height = 33
put(ms, 'A32', 'Amount caption (derived)', bold=True, h='right')
put(ms, 'B32', '=IF(B31=100,"(Amount in ₹ \'00)",IF(B31=1000,"(Amount in ₹ \'000)",IF(B31=1,"(Amount in ₹)","(Amount in ₹ / "&B31&")")))', h='left')
put(ms, 'A33', 'Current year column heading (derived)', bold=True, h='right')
put(ms, 'B33', '=UPPER(TEXT(B8,"mmmm d, yyyy"))', h='left')
put(ms, 'A34', 'Previous year column heading (derived)', bold=True, h='right')
put(ms, 'B34', '=UPPER(TEXT(B9,"mmmm d, yyyy"))', h='left')

# Partners' capital table
put(ms, 'A36', "PARTNERS' CAPITAL DETAILS (₹, as per Trial Balance)", bold=True, h='left')
put(ms, 'A37', 'Partner', bold=True); put(ms, 'B37', 'Capital Account ledger (Tally)', bold=True)
put(ms, 'C37', 'Contribution 31-Mar-26 (₹)', bold=True, h='center', wrap=True)
put(ms, 'D37', 'Contribution 31-Mar-25 (₹)', bold=True, h='center', wrap=True)
put(ms, 'E37', 'Profit sharing ratio (LLP Agreement)', bold=True, h='center', wrap=True)
ms.row_dimensions[37].height = 33
put(ms, 'A38', '=B22'); put(ms, 'B38', 'Samartha  R .N -Captial Account')
put(ms, 'C38', '=SUMIFS(TB!$D:$D,TB!$I:$I,"Partners Capital - Samartha R.N")', nf=NF_TB2)
put(ms, 'D38', '=SUMIFS(TB!$F:$F,TB!$I:$I,"Partners Capital - Samartha R.N")', nf=NF_TB2)
put(ms, 'E38', TBP, fill=YELLOW, h='center')
put(ms, 'A39', '=B26'); put(ms, 'B39', 'No capital account in Trial Balance', italic=True)
put(ms, 'C39', 0, nf=NF_TB2); put(ms, 'D39', 0, nf=NF_TB2); put(ms, 'E39', TBP, fill=YELLOW, h='center')
put(ms, 'A40', 'Total', bold=True); put(ms, 'C40', '=SUM(C38:C39)', bold=True, nf=NF_TB2); put(ms, 'D40', '=SUM(D38:D39)', bold=True, nf=NF_TB2)
box(ms, 'A37:E40')
put(ms, 'A42', 'Legend', bold=True, h='left')
put(ms, 'B42', 'Yellow cells = information not available in the Trial Balance; to be filled in before finalisation. All other cells are as per the Tally Trial Balances or derived by formula.', wrap=True, h='left')
ms.row_dimensions[42].height = 64
ms.merge_cells('B42:E42')
setup_print(ms, 'Master!$A$1:$E$42')

# =====================================================================================
# TB  (mapping engine)
# =====================================================================================
tb = wb.create_sheet('TB')
tb.sheet_view.showGridLines = False
widths(tb, {'A': 46, 'B': 40, 'C': 14.5, 'D': 14.5, 'E': 14.5, 'F': 14.5, 'G': 13, 'H': 13, 'I': 44, 'J': 30})
put(tb, 'A1', '=Master!B2', bold=True, size=10)
put(tb, 'A2', 'Trial Balance (closing balances as per Tally) - mapping to Financial Statements', bold=True)
put(tb, 'A3', '="Current year: "&Master!B10&"   |   Previous year: "&Master!B11', bold=True, size=10)
TOTALS_PLACEHOLDER = True
tb.merge_cells('C5:D5'); tb.merge_cells('E5:F5'); tb.merge_cells('G5:H5')
put(tb, 'C5', '2025-26 (₹) - as per Tally', bold=True, h='center', bottom=THIN)
put(tb, 'E5', '2024-25 (₹) - as per Tally', bold=True, h='center', bottom=THIN)
put(tb, 'G5', 'Rounded Off (÷ Master!B31)', bold=True, h='center', bottom=THIN, fill=YELLOW)
for col in 'DFH':
    tb[f'{col}5'].border = Border(bottom=THIN)
hdr = ['Particulars (Tally ledger)', 'Tally Group', 'Debit', 'Credit', 'Debit', 'Credit', 'Amount CY', 'Amount PY', 'Grouping (key used by Notes)', 'Head (Financial Statement caption)']
for i, hname in enumerate(hdr):
    put(tb, f'{get_column_letter(i+1)}6', hname, bold=True, size=10, color='FFFFFFFF', fill=BLUEHDR, h='center', v='center', wrap=True,
        top=THIN, bottom=THIN, left=THIN, right=THIN)
tb.row_dimensions[6].height = 30

# (ledger, group, CY dr, CY cr, PY dr, PY cr, key, head)
PC, LTB, OCL, OCA, TP, STLA, CCE, REV, OE, RS = ("Partners' capital", 'Long-term borrowings', 'Other current liabilities',
    'Other current assets', 'Trade payables', 'Short-term loans and advances', 'Cash and cash equivalents',
    'Revenue from operations', 'Other expenses', 'Reserves and surplus')
rows = [
 ('Samartha  R .N -Captial Account', 'Capital Account / Partners Capital', None, 90000, None, 90000, 'Partners Capital - Samartha R.N', PC),
 ('Manasa', 'Loans (Liability) / Unsecured Loans', None, 170180, None, None, 'Loan - Manasa', LTB),
 ('Samartha  Raghava NAgabhushanam-Loan', 'Loans (Liability) / Unsecured Loans', None, 2480680, None, 780680, 'Loan - Samartha Raghava Nagabhushanam', LTB),
 ('194 J-TDS on Profession', 'Current Liabilities / Duties & Taxes', None, 21760, None, 35000, 'TDS Payable 194J', OCL),
 ('GST Payable', 'Current Liabilities / Duties & Taxes', None, 7505.41, None, None, 'GST Payable', OCL),
 ('Input 9 % SGST', 'Current Liabilities / Duties & Taxes', 2820.94, None, None, None, 'Input SGST', OCA),
 ('Input  CGST 9%', 'Current Liabilities / Duties & Taxes', 2820.94, None, None, None, 'Input CGST', OCA),
]
creditors_cy = [('Ajay Pandey',30800),('Bhadri Narayan',8606.5),('Bharti Airtel Limited',6339.78),('Dhakshayani',3000),('Dhanush',8815),
 ('Face  Book India Online Services  Private Limited',163951.21),('Google India Private Limited',3657.53),('J R Computers Inc',9263),
 ('Karnataka Digital Studio',5400),('K Soorya',33930),('Madhavi K',90000),('Manjunath N',5000),('Pixel Digital and Design Studios',2000.9),
 ('Razorpay Software Pvt. Ltd.',30452.56),('Sathyanarayana Bhat',64590),('Sireesha',520),('Sragdhara Bhatt',2484),('Umesh',21500),
 ('Usha Devi',22500),('Venkatesh',16528),('Vinyas Kumar',18982)]
creditors_py = {'B Omkarmurthy':3000,'Medha Sudarshan':90000,'Sudhakar':67500,'Venkatesh':67500}
for n, amt in creditors_cy:
    rows.append((n, 'Current Liabilities / Sundry Creditors', None, amt, None, creditors_py.pop(n, None), 'Trade Payables - Others', TP))
for n, amt in creditors_py.items():
    rows.append((n, 'Current Liabilities / Sundry Creditors', None, None, None, amt, 'Trade Payables - Others', TP))
for n, amt in [('Mangala',15000),('Sriguru P V',20000),('Subodh Kumar Mishra',121000),('Swami Gangaram',2500),('Yuvaraj',15000)]:
    rows.append((n, 'Current Liabilities / Sundry Creditors (Dr balance)', amt, None, None, None, f'Advance - {n}', STLA))
rows += [
 ('Cash', 'Current Assets / Cash-in-Hand', 20754.2, None, None, None, 'Cash in hand', CCE),
 ('Kotak Mahindra Bank (Dr balance)', 'Current Assets / Bank Accounts', None, None, 12768.48, None, 'Bank - Kotak Mahindra Bank', CCE),
 ('Kotak Mahindra Bank (Cr balance - book overdraft)', 'Current Assets / Bank Accounts', None, 2593.05, None, None, 'Book overdraft - Kotak Mahindra Bank', OCL),
 ('B2C Sales', 'Sales Accounts', None, 1059363.54, None, None, 'B2C Sales', REV),
]
exp_cy = {'Accomodation and Hotel Expenses':17510,'Advertisement Campaining Cost IGST':410685.35,'Bank Charges':880.63,'Branding Expenses':19100,
'Conveyance':25373,'Data Infrastructure (GST)':26223.64,'Design and Development':70000,'GST Late File  Fees':20,'Interest on TDS':2164,
'Internship Charges':12000,'Local Travel and Conveyance':177417,'Mobile  and Telepohone Exenses':9871.95,'Office Expenses':80518,
'Payment Gate Way Comm. GST Exmt':2535.68,'Payment Gate Way Commission':24901.09,'Pooja Expenses':959127.5,'Postage and Courier Expenses':24538,
'Preincorporation Expenses':36000,'Printing Charges':43036,'Professional Charges':1008252,'Round Off':0.04,'Staff Welfare':7681,
'Studio Expenses (GST 0%)':24000,'Subscrptions Expenses':15000,'Translation Expenses':52760,'Web Site Development Expenses':10000}
exp_py = {'Bank Charges':311.52,'Employer Professional Tax':2500,'Interest on TDS':780,'Preincorporation Expenses':479320,'Professional Charges':450000,
'R O C and Filing  Expenses':35000,'Studio Expenses (GST 0%)':28000,'Web Site Development Expenses':125000}
# presentation names (spelling corrected) keyed by Tally ledger name
pretty = {'Accomodation and Hotel Expenses':'Accommodation and Hotel Expenses','Advertisement Campaining Cost IGST':'Advertisement Campaigning Cost',
 'GST Late File  Fees':'GST Late Filing Fees','Mobile  and Telepohone Exenses':'Mobile and Telephone Expenses','Payment Gate Way Comm. GST Exmt':'Payment Gateway Commission (GST Exempt)',
 'Payment Gate Way Commission':'Payment Gateway Commission','Preincorporation Expenses':'Pre-incorporation Expenses','Subscrptions Expenses':'Subscription Expenses',
 'Web Site Development Expenses':'Website Development Expenses','R O C and Filing  Expenses':'ROC and Filing Expenses','Studio Expenses (GST 0%)':'Studio Expenses'}
all_exp = sorted(set(exp_cy) | set(exp_py), key=lambda s: s.lower())
for n in all_exp:
    rows.append((n, 'Indirect Expenses', exp_cy.get(n), None, exp_py.get(n), None, n, OE))
rows.append(('Profit & Loss A/c (opening balance - accumulated loss brought forward)', 'Profit & Loss A/c', 1120911.52, None, None, None, 'Opening P&L', RS))

r = 7
for led, grp, cdr, ccr, pdr, pcr, key, head in rows:
    put(tb, f'A{r}', led, size=10); put(tb, f'B{r}', grp, size=10)
    for col, val in zip('CDEF', (cdr, ccr, pdr, pcr)):
        put(tb, f'{col}{r}', val, size=10, nf=NF_TB2)
    put(tb, f'G{r}', f'=(C{r}+D{r})/Master!$B$31', size=10, nf=NF_TB2)
    put(tb, f'H{r}', f'=(E{r}+F{r})/Master!$B$31', size=10, nf=NF_TB2)
    put(tb, f'I{r}', key, size=10); put(tb, f'J{r}', head, size=10)
    r += 1
last_tb = r - 1
box(tb, f'A7:J{last_tb}')
for col in 'CDEF':
    put(tb, f'{col}4', f'=SUM({col}7:{col}{last_tb+1})', size=10, nf=NF_TB2, bold=True)
put(tb, 'G4', f'=SUM(G7:G{last_tb+1})', size=10, nf=NF_TB, bold=True); put(tb, 'H4', f'=SUM(H7:H{last_tb+1})', size=10, nf=NF_TB, bold=True)
put(tb, 'A4', f'Totals of rows 7 to {last_tb} (insert new ledger rows above row {last_tb+1})', size=9, italic=True, color='FF808080')
r += 1
put(tb, f'A{r}', 'Grand Total as per Tally Trial Balance', bold=True, size=10)
put(tb, f'C{r}', 4380402.48, size=10, nf=NF_TB2); put(tb, f'D{r}', 4380402.48, size=10, nf=NF_TB2)
put(tb, f'E{r}', 1133680, size=10, nf=NF_TB2); put(tb, f'F{r}', 1133680, size=10, nf=NF_TB2)
r += 1
put(tb, f'A{r}', 'Difference: mapped total less Tally Grand Total (must be nil)', bold=True, size=10)
for col in 'CDEF':
    put(tb, f'{col}{r}', f'={col}4-{col}{r-1}', size=10, nf=NF_TB2, bold=True)
r += 1
put(tb, f'A{r}', 'Difference: Debit less Credit (must be nil)', bold=True, size=10)
put(tb, f'C{r}', '=C4-D4', size=10, nf=NF_TB2, bold=True); put(tb, f'E{r}', '=E4-F4', size=10, nf=NF_TB2, bold=True)
TB_CHECK_ROWS = (r-1, r)
tb.freeze_panes = 'A7'
setup_print(tb, f'TB!$A$1:$J${r}', landscape=True)

# helpers for notes formulas
def cy(key_ref):  return f'=SUMIFS(TB!$G:$G,TB!$I:$I,{key_ref})'
def py(key_ref):  return f'=SUMIFS(TB!$H:$H,TB!$I:$I,{key_ref})'

# =====================================================================================
# common header block for statements
# =====================================================================================
def header_block(ws, title_formula, first_col='A', caption_col='E', note_title=False):
    fc = first_col
    put(ws, f'{fc}1', '=UPPER(Master!B2)', bold=True)
    put(ws, f'{fc}2', '="LLPIN: "&UPPER(Master!B4)', bold=True)
    put(ws, f'{fc}3', '=Master!B3', bold=True, v='top')
    ws.row_dimensions[3].height = 33.75
    put(ws, f'{fc}4', title_formula, bold=True)
    put(ws, f'{caption_col}5', '=Master!B32', bold=not note_title, h='right')
    ws.row_dimensions[5].height = 15

def sign_block(ws, start_row, cols=('A', 'B', 'D'), indent_b=8, wrap_a=True):
    a, b, d = cols
    r = start_row
    put(ws, f'{a}{r}', 'As per our report attached', wrap=wrap_a)
    put(ws, f'{b}{r}', 'For and on behalf of the Partners of', h='left', indent=indent_b)
    put(ws, f'{a}{r+1}', '="For "&Master!B16', bold=True)
    put(ws, f'{b}{r+1}', '=Master!B2', bold=True, h='left', indent=indent_b)
    put(ws, f'{a}{r+2}', '="(FRN: "&Master!B17&")"', italic=True)
    r += 6
    put(ws, f'{a}{r}', '=UPPER(Master!B13)', bold=True, h='left', v='center', wrap=wrap_a)
    put(ws, f'{b}{r}', '=UPPER(Master!B22)', bold=True, h='left', v='top')
    put(ws, f'{d}{r}', '=UPPER(Master!B26)', bold=True, h='left', v='top')
    put(ws, f'{a}{r+1}', '=PROPER(Master!B14)', italic=True)
    put(ws, f'{b}{r+1}', 'Designated Partner', italic=True, h='left')
    put(ws, f'{d}{r+1}', f'={b}{r+1}', italic=True, h='left')
    put(ws, f'{a}{r+2}', '="Membership No. "&Master!B15')
    put(ws, f'{b}{r+2}', '="DPIN: "&Master!B23', h='left')
    put(ws, f'{d}{r+2}', '="DPIN: "&Master!B27', h='left')
    put(ws, f'{a}{r+3}', '="UDIN: "&Master!B18')
    put(ws, f'{a}{r+4}', '=Master!B19', h='left'); put(ws, f'{b}{r+4}', '=Master!B24', h='left'); put(ws, f'{d}{r+4}', '=Master!B28', h='left')
    put(ws, f'{a}{r+5}', '=Master!B20', h='left'); put(ws, f'{b}{r+5}', '=Master!B25', h='left'); put(ws, f'{d}{r+5}', '=Master!B29', h='left')
    return r + 5

def col_headers(ws, row, cols=('A','B','C','D','E'), note_col='C', cy_col='D', py_col='E', label='PARTICULARS', label_col='A'):
    for col in cols:
        put(ws, f'{col}{row}', None, bold=True, h='center', v='center', top=MED, bottom=MED)
    ws[f'{label_col}{row}'].value = label
    if note_col: ws[f'{note_col}{row}'].value = 'NOTE'
    put(ws, f'{cy_col}{row}', '=Master!B33', bold=True, h='right', v='center', top=MED, bottom=MED)
    put(ws, f'{py_col}{row}', '=Master!B34', bold=True, h='right', v='center', top=MED, bottom=MED)

# =====================================================================================
# BS
# =====================================================================================
bs = wb.create_sheet('BS')
widths(bs, {'A': 50.0, 'B': 32.0, 'C': 6.4, 'D': 20.0, 'E': 20.0, 'F': 4, 'G': 30})
header_block(bs, '="BALANCE SHEET AS AT "&Master!B33')
col_headers(bs, 6)
bs.row_dimensions[6].height = 17.25
COLS = ('A','B','C','D','E')
def line(ws, r, text, note=None, cyf=None, pyf=None, bold=False, indent=0, amt_bold=False, top=None, bottom=None, h='left', hdr_border=False):
    put(ws, f'A{r}', text, bold=bold, h=h, wrap=True, indent=indent, nf='@')
    if note is not None: put(ws, f'C{r}', note, h='center', nf='@')
    put(ws, f'D{r}', cyf, bold=amt_bold or bold, nf=NF_AMT)
    put(ws, f'E{r}', pyf, bold=amt_bold or bold, nf=NF_AMT)
    if top or bottom:
        for col in ('D', 'E'):
            ws[f'{col}{r}'].border = Border(top=top or NONE, bottom=bottom or NONE)
    ws.row_dimensions[r].height = 17.1

line(bs, 7, 'LIABILITIES', bold=True); row_border(bs, 7, COLS, top=MED)
line(bs, 8, "Partners' funds", bold=True)
line(bs, 9, "(a) Partners' capital", '2.1', "='2.1'!D13", "='2.1'!E13", indent=1)
line(bs, 10, '(b) Reserves and surplus', '2.2', "='2.2 - 2.10'!D15", "='2.2 - 2.10'!E15", indent=1)
line(bs, 11, "Total Partners' funds", None, '=SUM(D9:D10)', '=SUM(E9:E10)', indent=2, amt_bold=True, top=THIN, bottom=THIN)
line(bs, 12, 'Non-current liabilities', bold=True)
line(bs, 13, '(a) Long-term borrowings', '2.3', "='2.2 - 2.10'!D22", "='2.2 - 2.10'!E22", indent=1, bottom=THIN)
line(bs, 14, 'Total Non-current liabilities', None, '=SUM(D13:D13)', '=SUM(E13:E13)', indent=2, amt_bold=True, top=THIN, bottom=THIN)
line(bs, 15, 'Current Liabilities', bold=True)
line(bs, 16, '(a) Trade payables', '2.4', "='2.2 - 2.10'!D30", "='2.2 - 2.10'!E30", indent=1)
line(bs, 17, '(b) Other current liabilities', '2.5', "='2.2 - 2.10'!D39", "='2.2 - 2.10'!E39", indent=1, bottom=THIN)
line(bs, 18, 'Total Current Liabilities', None, '=SUM(D16:D17)', '=SUM(E16:E17)', indent=2, amt_bold=True, top=THIN, bottom=MED)
line(bs, 19, 'TOTAL LIABILITIES', None, '=D11+D14+D18', '=E11+E14+E18', bold=True)
row_border(bs, 19, COLS, top=MED, bottom=MED)
for col in ('D','E'): bs[f'{col}19'].alignment = Alignment(vertical='center')
bs.row_dimensions[20].height = 5.1
line(bs, 21, 'ASSETS', bold=True)
line(bs, 22, 'Non-current assets', bold=True)
line(bs, 23, '(a) Property, Plant & Equipment & Intangible assets', indent=1)
line(bs, 24, '(i) Property, Plant & Equipment', None, 0, 0, indent=2)
line(bs, 25, '="Total "&A22', None, '=SUM(D24:D24)', '=SUM(E24:E24)', indent=2, amt_bold=True, top=THIN, bottom=THIN)
line(bs, 26, 'Current assets', bold=True)
line(bs, 27, '(a) Cash and cash equivalents', '2.6', "='2.2 - 2.10'!D45", "='2.2 - 2.10'!E45", indent=1)
line(bs, 28, '(b) Short-term loans and advances', '2.7', "='2.2 - 2.10'!D55", "='2.2 - 2.10'!E55", indent=1)
line(bs, 29, '(c) Other current assets', '2.8', "='2.2 - 2.10'!D62", "='2.2 - 2.10'!E62", indent=1)
line(bs, 30, '="Total "&A26', None, '=SUM(D27:D29)', '=SUM(E27:E29)', indent=2, amt_bold=True, top=THIN, bottom=MED)
line(bs, 31, '=UPPER("Total "&A21)', None, '=D25+D30', '=E25+E30', bold=True)
row_border(bs, 31, COLS, top=MED, bottom=MED)
bs.row_dimensions[32].height = 5.1
put(bs, 'A33', 'SIGNIFICANT ACCOUNTING POLICIES,\nNOTES FORMING PART OF THE FINANCIAL STATEMENTS', wrap=True, nf='@')
put(bs, 'C33', '1 & 2', h='center'); bs.row_dimensions[33].height = 30
sign_end = sign_block(bs, 35)
# off-print check cells
put(bs, 'G6', 'Check (outside print area)', bold=True, size=9, color='FF808080')
put(bs, 'G7', '="Total Liabilities less Total Assets (CY): "&TEXT(D19-D31,"#,##0.00")', size=9, color='FF808080')
put(bs, 'G8', '="Total Liabilities less Total Assets (PY): "&TEXT(E19-E31,"#,##0.00")', size=9, color='FF808080')
setup_print(bs, f'BS!$A$1:$E${sign_end}')

# =====================================================================================
# PL
# =====================================================================================
pl = wb.create_sheet('PL')
widths(pl, {'A': 61.5, 'B': 32.0, 'C': 6.4, 'D': 20.0, 'E': 20.0})
header_block(pl, '="PROFIT AND LOSS STATEMENT FOR THE PERIOD "&Master!B5')
col_headers(pl, 6); pl.row_dimensions[6].height = 17.45
line(pl, 7, 'REVENUE', bold=True); row_border(pl, 7, ('D','E'), top=MED)
line(pl, 8, '(a) Revenue from operations', '2.9', "='2.2 - 2.10'!D68", "='2.2 - 2.10'!E68", indent=1)
line(pl, 9, '(b) Other income', None, 0, 0, indent=1, bottom=THIN)
line(pl, 10, 'Total Income', None, '=SUM(D8:D9)', '=SUM(E8:E9)', bold=True, indent=2, top=THIN, bottom=THIN)
pl.row_dimensions[11].height = 17.1
line(pl, 12, 'EXPENSES', bold=True)
line(pl, 13, '(A) Other expenses', '2.10', "='2.2 - 2.10'!D105", "='2.2 - 2.10'!E105", indent=1, bottom=THIN)
line(pl, 14, 'Total Expenses', None, '=SUM(D13:D13)', '=SUM(E13:E13)', bold=True, indent=2, top=THIN, bottom=THIN)
pl.row_dimensions[15].height = 17.1
line(pl, 16, 'Profit/(Loss) before exceptional & extraordinary items and tax', None, '=D10-D14', '=E10-E14', bold=True, top=THIN, bottom=THIN)
line(pl, 17, 'Exceptional items', None, 0, 0, top=THIN, bottom=THIN)
line(pl, 18, 'Profit/(Loss) before extraordinary items and tax', None, '=D16-D17', '=E16-E17', bold=True, top=THIN, bottom=THIN)
line(pl, 19, 'Extraordinary items', None, 0, 0, top=THIN, bottom=THIN)
line(pl, 20, 'Profit/(Loss) before tax', None, '=D18+D19', '=E18+E19', bold=True, top=THIN, bottom=THIN)
pl.row_dimensions[21].height = 5.1
line(pl, 22, 'Tax Expenses', bold=True)
line(pl, 23, '(a) Current Tax', None, 0, 0, indent=2)
line(pl, 24, '(b) Deferred Tax Charge/(Credit)', None, 0, 0, indent=2, bottom=THIN)
line(pl, 25, 'Total tax expenses', None, '=SUM(D23:D24)', '=SUM(E23:E24)', bold=True, indent=2, top=THIN, bottom=THIN)
pl.row_dimensions[26].height = 5.1
line(pl, 27, 'Profit/(Loss) after tax', None, '=D20-D25', '=E20-E25', bold=True)
row_border(pl, 27, COLS, top=MED, bottom=MED)
pl.row_dimensions[28].height = 5.1
put(pl, 'A29', 'SIGNIFICANT ACCOUNTING POLICIES,\nNOTES FORMING PART OF THE FINANCIAL STATEMENTS', wrap=True, nf='@')
put(pl, 'C29', '1 & 2', h='center'); pl.row_dimensions[29].height = 33
sign_end_pl = sign_block(pl, 31, indent_b=3)
setup_print(pl, f'PL!$A$1:$E${sign_end_pl}')

# =====================================================================================
# 2.1 Partners' capital
# =====================================================================================
n1 = wb.create_sheet('2.1')
widths(n1, {'A': 46.1, 'B': 18.5, 'C': 16.6, 'D': 17.5, 'E': 17.5, 'F': 44})
header_block(n1, '="2. NOTES FORMING PART OF THE FINANCIAL STATEMENTS FOR THE YEAR ENDED "&Master!B33', caption_col='E', note_title=True)
n1['E5'].value = None
put(n1, 'A6', "2.1             Partners' capital", bold=True, top=THIN, bottom=THIN, left=THIN, right=THIN)
n1.row_dimensions[7].height = 5.1
put(n1, 'A8', "(A) Statement Showing Partners' Capital Contribution", bold=True, h='left', indent=1)
put(n1, 'E8', '=Master!B32', h='right')
col_headers(n1, 9, note_col=None)
put(n1, 'A10', 'Capital contribution by partners', bold=True, size=10)
# keys in column F (outside print area)
put(n1, 'F11', 'Partners Capital - Samartha R.N', size=9, color='FF808080')
put(n1, 'A11', '=Master!B22', size=10, h='left', indent=1)
put(n1, 'D11', cy('$F11'), nf=NF_AMT); put(n1, 'E11', py('$F11'), nf=NF_AMT)
put(n1, 'A12', '=Master!B26', size=10, h='left', indent=1)
put(n1, 'D12', 0, nf=NF_AMT); put(n1, 'E12', 0, nf=NF_AMT)
n1.row_dimensions[12].hidden = False
put(n1, 'A13', 'Total', bold=True, size=10, h='left', indent=1)
put(n1, 'D13', '=SUM(D11:D12)', bold=True, nf=NF_AMT, top=THIN, bottom=THIN); put(n1, 'E13', '=SUM(E11:E12)', bold=True, nf=NF_AMT, top=THIN, bottom=THIN)
n1.row_dimensions[14].height = 5.1
put(n1, 'A15', "(B) Reconciliation Of Partners' Capital At The Beginning And At The End Of The Reporting Period", bold=True, h='left', indent=1)
put(n1, 'E15', '=Master!B32', h='right')
col_headers(n1, 16, note_col=None)
put(n1, 'A17', 'Balance at the beginning of the year'); put(n1, 'D17', '=E19', nf=NF_AMT); put(n1, 'E17', 0, nf=NF_AMT)
put(n1, 'A18', 'Add: Capital contributed during the year'); put(n1, 'D18', '=D19-D17', nf=NF_AMT); put(n1, 'E18', '=E19-E17', nf=NF_AMT)
put(n1, 'A19', 'Balance at the end of the year', top=THIN, bottom=THIN)
for col in 'BC': put(n1, f'{col}19', None, top=THIN, bottom=THIN)
put(n1, 'D19', '=D13', nf=NF_AMT, top=THIN, bottom=THIN); put(n1, 'E19', '=E13', nf=NF_AMT, top=THIN, bottom=THIN)
n1.row_dimensions[20].height = 5.1
put(n1, 'A21', '(C) Terms Of Partnership:', bold=True, h='left', indent=1)
n1.merge_cells('A22:E22')
put(n1, 'A22', "The rights and obligations of the partners, their capital contribution and the profit sharing ratio are governed by the LLP Agreement. "
    "No interest on partners' capital and no remuneration to partners has been charged to the Statement of Profit and Loss for the year "
    "(previous year: Nil), as no such ledgers appear in the books of account.", wrap=True, v='top')
n1.row_dimensions[22].height = 52
n1.row_dimensions[23].height = 5.1
put(n1, 'A24', "(D) Details Of Partners And Their Capital Contribution:", bold=True, h='left', indent=1)
for col in ('A','B','C','D','E'):
    put(n1, f'{col}25', None, bold=True, h='center', v='center', top=MED, bottom=THIN)
n1['A25'].value = 'PARTICULARS'
put(n1, 'D25', '=Master!B33', bold=True, h='right', v='center', top=MED, bottom=THIN)
put(n1, 'E25', '=Master!B34', bold=True, h='right', v='center', top=MED, bottom=THIN)
for col in ('A','B','C'):
    put(n1, f'{col}26', None, bottom=MED)
put(n1, 'D26', 'Contribution', bold=True, h='right', bottom=MED); put(n1, 'E26', 'Contribution', bold=True, h='right', bottom=MED)
put(n1, 'A27', '=A11'); put(n1, 'D27', '=D11', nf=NF_AMT); put(n1, 'E27', '=E11', nf=NF_AMT)
put(n1, 'A28', '=A12'); put(n1, 'D28', '=D12', nf=NF_AMT); put(n1, 'E28', '=E12', nf=NF_AMT)
put(n1, 'A29', 'Total', bold=True, top=THIN, bottom=THIN)
for col in 'BC': put(n1, f'{col}29', None, top=THIN, bottom=THIN)
put(n1, 'D29', '=SUM(D27:D28)', bold=True, nf=NF_AMT, top=THIN, bottom=THIN); put(n1, 'E29', '=SUM(E27:E28)', bold=True, nf=NF_AMT, top=THIN, bottom=THIN)
put(n1, 'A30', '% of total contribution', italic=True)
put(n1, 'D30', '=IF(D29=0,0,D27/D29)', nf='0%', h='right'); put(n1, 'E30', '=IF(E29=0,0,E27/E29)', nf='0%', h='right')
put(n1, 'F1', 'Grouping keys (outside print area)', size=9, color='FF808080')
setup_print(n1, "'2.1'!$A$1:$E$30")

# =====================================================================================
# 2.2 - 2.10
# =====================================================================================
n2 = wb.create_sheet('2.2 - 2.10')
widths(n2, {'A': 40.0, 'B': 11.9, 'C': 70.6, 'D': 17.4, 'E': 17.4})
header_block(n2, '="2. NOTES FORMING PART OF THE FINANCIAL STATEMENTS FOR THE YEAR ENDED "&Master!B33', first_col='B', caption_col='E', note_title=True)
put(n2, 'A1', 'Grouping keys (column A is outside the print area)', size=9, color='FF808080')
for col in ('B','C','D','E'):
    put(n2, f'{col}6', None, bold=True, h='center', v='center', top=MED, bottom=MED)
n2['B6'].value = 'NOTE'; n2['C6'].value = 'PARTICULARS'
put(n2, 'D6', '=Master!B33', bold=True, h='center', v='center', wrap=True, top=MED, bottom=MED)
put(n2, 'E6', '=Master!B34', bold=True, h='center', v='center', wrap=True, top=MED, bottom=MED)
n2.row_dimensions[6].height = 17.25; n2.row_dimensions[7].height = 5.1

def note_title(ws, r, num, title):
    put(ws, f'B{r}', num, bold=True, h='center', nf='@', top=THIN, bottom=THIN, left=THIN)
    put(ws, f'C{r}', title, bold=True, h='left', v='center', top=THIN, bottom=THIN, right=THIN)
def n_item(ws, r, text, key=None, indent=1, cyv=None, pyv=None, bold=False, top=None, bottom=None, size=11):
    if key is not None:
        put(ws, f'A{r}', key, size=9, color='FF808080')
        cyv, pyv = cy(f'$A{r}'), py(f'$A{r}')
    put(ws, f'C{r}', text, h='left', indent=indent, bold=bold, wrap=True, size=size)
    put(ws, f'D{r}', cyv, nf=NF_AMT, bold=bold, top=top, bottom=bottom)
    put(ws, f'E{r}', pyv, nf=NF_AMT, bold=bold, top=top, bottom=bottom)
def n_total(ws, r, text, rng_from, rng_to, bottom=THIN, formula=None):
    put(ws, f'C{r}', text, bold=True, h='left', indent=1)
    for col in ('D','E'):
        f = formula.replace('#', col) if formula else f'=SUM({col}{rng_from}:{col}{rng_to})'
        put(ws, f'{col}{r}', f, bold=True, nf=NF_AMT, top=THIN, bottom=bottom)

# 2.2 Reserves and surplus  (rows 8-16)
note_title(n2, 8, '2.2', 'Reserves and surplus'); n2.row_dimensions[9].height = 5.1
put(n2, 'C10', 'Surplus / (deficit) - balance in the statement of profit and loss account', bold=True, top=THIN)
for col in ('B','D','E'): put(n2, f'{col}10', None, top=THIN)
put(n2, 'A11', 'Opening P&L', size=9, color='FF808080')
n_item(n2, 11, 'Opening Balance', cyv='=-SUMIFS(TB!$G:$G,TB!$I:$I,$A11)', pyv='=-SUMIFS(TB!$H:$H,TB!$I:$I,$A11)')
n_item(n2, 12, 'Add: Profit/ (Loss) for the year', cyv='=PL!D27', pyv='=PL!E27')
n_item(n2, 13, 'Less: Amount utilised', cyv=0, pyv=0)
n_item(n2, 14, 'Less: Transfer to reserves', cyv=0, pyv=0)
n_total(n2, 15, 'Closing Balance', 11, 14, formula='=#11+#12-#13-#14')
put(n2, 'D16', '=D15', bold=True, nf=NF_AMT, top=THIN, bottom=THIN); put(n2, 'E16', '=E15', bold=True, nf=NF_AMT, top=THIN, bottom=THIN)
for col in ('B','C'): put(n2, f'{col}16', None, bottom=THIN)
n2.row_dimensions[17].height = 5.1

# 2.3 Long-term borrowings (rows 18-22)
note_title(n2, 18, '2.3', 'Long-term borrowings')
put(n2, 'C19', 'Unsecured loans', bold=True, top=THIN)
for col in ('B','D','E'): put(n2, f'{col}19', None, top=THIN)
n_item(n2, 20, '="From partner - "&Master!B22', key='Loan - Samartha Raghava Nagabhushanam')
n_item(n2, 21, 'From others - Manasa', key='Loan - Manasa')
n_total(n2, 22, 'Total', 20, 21)
n2.row_dimensions[23].height = 5.1

# 2.4 Trade payables (rows 24-30)
note_title(n2, 24, '2.4', 'Trade payables')
n2.row_dimensions[25].height = 5.1
put(n2, 'C26', 'Outstanding dues of;', top=THIN)
for col in ('B','D','E'): put(n2, f'{col}26', None, top=THIN)
n_item(n2, 27, '(a) micro enterprises and small enterprises', key='Trade Payables - MSME', indent=0)
n_item(n2, 28, '(b) other than micro enterprises and small enterprises', key='Trade Payables - Others', indent=0)
n2.row_dimensions[29].height = 5.1
n_total(n2, 30, 'Total', 27, 28)
n2.row_dimensions[31].height = 5.1

# 2.5 Other current liabilities (rows 32-39)
note_title(n2, 32, '2.5', 'Other current liabilities')
put(n2, 'C33', 'Statutory dues', bold=True, top=THIN)
for col in ('B','D','E'): put(n2, f'{col}33', None, top=THIN)
n_item(n2, 34, 'TDS payable (Section 194J)', key='TDS Payable 194J')
n_item(n2, 35, 'GST payable', key='GST Payable')
put(n2, 'C36', 'Others', bold=True)
n_item(n2, 37, 'Book overdraft - Kotak Mahindra Bank (credit balance as per books)', key='Book overdraft - Kotak Mahindra Bank')
n2.row_dimensions[38].height = 5.1
n_total(n2, 39, 'Total', 34, 37)
n2.row_dimensions[40].height = 5.1

# 2.6 Cash and cash equivalents (rows 41-45)
note_title(n2, 41, '2.6', 'Cash and cash equivalents')
put(n2, 'B42', None, top=THIN); put(n2, 'C42', None, top=THIN); put(n2, 'D42', None, top=THIN); put(n2, 'E42', None, top=THIN)
n2.row_dimensions[42].height = 5.1
n_item(n2, 43, 'Cash in hand', key='Cash in hand')
n_item(n2, 44, 'Balances with banks - Kotak Mahindra Bank (current account)', key='Bank - Kotak Mahindra Bank')
n_total(n2, 45, 'Total', 43, 44)
n2.row_dimensions[46].height = 5.1

# 2.7 Short-term loans and advances (rows 47-55)
note_title(n2, 47, '2.7', 'Short-term loans and advances')
put(n2, 'C48', 'Advances to suppliers / consultants (debit balances in Sundry Creditors)', bold=True, top=THIN)
for col in ('B','D','E'): put(n2, f'{col}48', None, top=THIN)
adv = ['Mangala', 'Sriguru P V', 'Subodh Kumar Mishra', 'Swami Gangaram', 'Yuvaraj']
for i, nm in enumerate(adv):
    n_item(n2, 49 + i, nm, key=f'Advance - {nm}')
n2.row_dimensions[54].height = 5.1
n_total(n2, 55, 'Total', 49, 53)
n2.row_dimensions[56].height = 5.1

# 2.8 Other current assets (rows 57-62)
note_title(n2, 57, '2.8', 'Other current assets')
put(n2, 'C58', 'Balances with government authorities', bold=True, top=THIN)
for col in ('B','D','E'): put(n2, f'{col}58', None, top=THIN)
n_item(n2, 59, 'Input CGST 9%', key='Input CGST')
n_item(n2, 60, 'Input SGST 9%', key='Input SGST')
n2.row_dimensions[61].height = 5.1
n_total(n2, 62, 'Total', 59, 60)
n2.row_dimensions[63].height = 5.1

# 2.9 Revenue from operations (rows 64-68)
note_title(n2, 64, '2.9', 'Revenue from operations')
put(n2, 'B65', None, top=THIN); put(n2, 'C65', None, top=THIN); put(n2, 'D65', None, top=THIN); put(n2, 'E65', None, top=THIN)
n2.row_dimensions[65].height = 5.1
n_item(n2, 66, 'B2C Sales', key='B2C Sales')
n2.row_dimensions[67].height = 5.1
n_total(n2, 68, 'Total', 66, 66)
n2.row_dimensions[69].height = 5.1

# 2.10 Other expenses (rows 70-105)
note_title(n2, 70, '2.10', 'Other expenses')
put(n2, 'B71', None, top=THIN); put(n2, 'C71', None, top=THIN); put(n2, 'D71', None, top=THIN); put(n2, 'E71', None, top=THIN)
n2.row_dimensions[71].height = 5.1
r = 72
for nm in all_exp:
    n_item(n2, r, pretty.get(nm, nm), key=nm)
    r += 1
exp_last = r - 1  # 72 + 28 - 1 = 99
r_total = 105
for rr in range(exp_last + 1, r_total):
    n2.row_dimensions[rr].height = 5.1
n_total(n2, r_total, 'Total', 72, exp_last, bottom=THIN)
n2.row_dimensions[r_total + 1].height = 5.1
assert exp_last < r_total, exp_last
setup_print(n2, f"'2.2 - 2.10'!$B$1:$E${r_total}")
from openpyxl.worksheet.pagebreak import Break
n2.row_breaks.append(Break(id=69))
n2.print_title_rows = '1:6'

# =====================================================================================
# 2.11 - 2.12
# =====================================================================================
n3 = wb.create_sheet('2.11-2.12')
widths(n3, {'A': 9.4, 'B': 46.0, 'C': 40.0, 'D': 26.0, 'E': 20.0})
put(n3, 'A1', "='2.2 - 2.10'!B1", bold=True); put(n3, 'A2', "='2.2 - 2.10'!B2", bold=True)
put(n3, 'A3', '=Master!B3', bold=True, v='top'); n3.row_dimensions[3].height = 35.25
put(n3, 'A4', "='2.2 - 2.10'!B4", bold=True)
put(n3, 'A6', '2.11  Related Party Disclosures', bold=True, h='left', v='top', top=THIN, bottom=THIN, left=THIN, right=THIN)
put(n3, 'B6', None, top=THIN, bottom=THIN, right=THIN)
n3.merge_cells('A6:B6')
n3.row_dimensions[7].height = 5.1
put(n3, 'A8', 'a. Designated Partners / Key Management Personnel:', bold=True)
n3.row_dimensions[9].height = 5.1
put(n3, 'A10', 'Sl No', bold=True, h='center', top=THIN, bottom=THIN)
put(n3, 'B10', 'Name of the Party', bold=True, h='left', top=THIN, bottom=THIN)
put(n3, 'C10', None, top=THIN, bottom=THIN)
put(n3, 'D10', 'Nature of relationship', bold=True, h='left', top=THIN, bottom=THIN)
put(n3, 'E10', None, top=THIN, bottom=THIN)
put(n3, 'A11', 1, h='center', v='center'); put(n3, 'B11', '=Master!B22'); put(n3, 'D11', 'Designated Partner', h='left')
put(n3, 'A12', '=+A11+1', h='center', v='center', bottom=THIN); put(n3, 'B12', '=Master!B26', bottom=THIN)
put(n3, 'C12', None, bottom=THIN); put(n3, 'D12', 'Designated Partner', h='left', bottom=THIN); put(n3, 'E12', None, bottom=THIN)
n3.row_dimensions[13].height = 5.1
put(n3, 'A14', 'b. Related Party Transactions and Balances:', bold=True, size=10, v='center')
put(n3, 'E14', '=Master!B32', bold=True, size=10, h='right')
n3.row_dimensions[15].height = 5.1
put(n3, 'A16', 'Sl.No.', bold=True, size=10, h='center', v='center', top=THIN)
put(n3, 'B16', 'Nature of transaction', bold=True, size=10, h='center', v='center', top=THIN)
put(n3, 'C16', None, top=THIN)
put(n3, 'D16', 'For the year ended', bold=True, size=10, h='right', v='top', wrap=True, top=THIN)
put(n3, 'E16', 'For the year ended', bold=True, size=10, h='right', v='top', wrap=True, top=THIN)
put(n3, 'A17', None, bottom=THIN); put(n3, 'B17', None, bottom=THIN); put(n3, 'C17', None, bottom=THIN)
put(n3, 'D17', '=PROPER(Master!B33)', bold=True, size=10, h='right', v='top', top=THIN, bottom=THIN)
put(n3, 'E17', '=PROPER(Master!B34)', bold=True, size=10, h='right', v='top', top=THIN, bottom=THIN)
put(n3, 'A18', 1, size=10, h='center'); put(n3, 'B18', "Partners' capital contribution (closing balance)", bold=True, size=10)
put(n3, 'B19', '=B11', size=10); put(n3, 'D19', "='2.1'!D11", size=10, nf=NF_TB, h='center', v='center'); put(n3, 'E19', "='2.1'!E11", size=10, nf=NF_TB, h='center', v='center')
n3.row_dimensions[20].height = 5.1
put(n3, 'A21', 2, size=10, h='center'); put(n3, 'B21', 'Unsecured loan from partner (closing balance)', bold=True, size=10)
put(n3, 'B22', '=B11', size=10); put(n3, 'D22', "='2.2 - 2.10'!D20", size=10, nf=NF_TB, h='center', v='center'); put(n3, 'E22', "='2.2 - 2.10'!E20", size=10, nf=NF_TB, h='center', v='center')
n3.row_dimensions[23].height = 5.1
put(n3, 'A24', 3, size=10, h='center'); put(n3, 'B24', 'Unsecured loan from partner - net amount received during the year', bold=True, size=10, wrap=True)
put(n3, 'B25', '=B11', size=10, bottom=THIN); put(n3, 'A25', None, bottom=THIN); put(n3, 'C25', None, bottom=THIN)
put(n3, 'D25', '=D22-E22', size=10, nf=NF_TB, h='center', v='center', bottom=THIN); put(n3, 'E25', '=E22-0', size=10, nf=NF_TB, h='center', v='center', bottom=THIN)
n3.row_dimensions[26].height = 5.1
n3.merge_cells('A27:E27')
put(n3, 'A27', 'No interest on the loan from the partner and no remuneration to partners has been charged in the books for the year (previous year: Nil). '
    'The previous-year figure of net loan received is stated on the basis that no opening balances appear in the Trial Balance for 2024-25.', wrap=True, v='top', h='justify', size=10)
n3.row_dimensions[27].height = 42
n3.row_dimensions[28].height = 5.1
put(n3, 'A29', '2.12 - Previous period figures', bold=True, top=THIN, bottom=THIN, left=THIN, right=THIN)
put(n3, 'B29', None, top=THIN, bottom=THIN, right=THIN)
n3.merge_cells('A29:B29')
n3.merge_cells('A30:E30')
put(n3, 'A30', "Previous year's figures have been regrouped / reclassified wherever necessary to conform to the current year's presentation.", wrap=True, h='justify', v='top')
n3.row_dimensions[30].height = 33
sign_end_n3 = sign_block(n3, 32, cols=('A', 'C', 'D'), indent_b=0, wrap_a=False)
setup_print(n3, f"'2.11-2.12'!$A$1:$E${sign_end_n3}")

# =====================================================================================
# FLAGS
# =====================================================================================
fl = wb.create_sheet('Flags')
fl.sheet_view.showGridLines = False
widths(fl, {'A': 5, 'B': 26, 'C': 70, 'D': 60})
put(fl, 'B1', 'ITEMS REQUIRING CLARIFICATION / CONFIRMATION - NOT ASSUMED IN THE FINANCIAL STATEMENTS', bold=True)
put(fl, 'B2', '=Master!B2&" - FY "&Master!B5', bold=True)
for i, h in enumerate(['#', 'Area', 'Observation (as per Trial Balance)', 'Information required / treatment adopted']):
    put(fl, f'{get_column_letter(i+1)}4', h, bold=True, color='FFFFFFFF', fill=BLUEHDR, h='center', v='center', wrap=True, top=THIN, bottom=THIN, left=THIN, right=THIN)
flags = [
 ('Period', 'The FY 2025-26 Trial Balance exported from Tally covers 1-Apr-25 to 23-Mar-26, not the full year to 31-Mar-26.',
  'Statements are presented as at March 31, 2026 (FY 2025-26) as instructed. Transactions from 24-Mar-26 to 31-Mar-26 (if any) are not included. A TB to 31-Mar-26 should be re-exported and pasted into the TB sheet.'),
 ('Entity / partners', 'Only one partner capital account appears in the TB (Samartha R .N - Captial Account, Rs 90,000). An LLP must have at least two partners.',
  'Second designated partner name, DPIN, capital (if any) and profit-sharing ratio are required. Placeholders are in the Master sheet (yellow cells). Note 2.1(D) currently shows 100% contribution by the one partner in the TB.'),
 ('Entity details', 'LLPIN, PAN, LLP Agreement terms, incorporation date, auditor name/firm/FRN/membership no./UDIN, signing dates and places are not available in the TB.',
  'Fill the yellow cells in the Master sheet; all headers and signature blocks are linked to them. Note: statutory audit of an LLP is mandatory only if turnover exceeds Rs 40 lakh or contribution exceeds Rs 25 lakh; neither threshold is met per the TB, so confirm whether an auditor block is required.'),
 ('Format', 'The reference (Checkaro) is a Private Limited Company under Schedule III. The entity here is an LLP.',
  'Layout, fonts, rounding and note structure follow the reference exactly. Captions have been changed only where the entity type dictates: Share capital -> Partners\' capital; Shareholder\'s funds -> Partners\' funds; Board of Directors/Director/DIN -> Partners/Designated Partner/DPIN. Earnings per share and promoter shareholding (not applicable to an LLP) are omitted.'),
 ('Bank', 'Kotak Mahindra Bank ledger has a CREDIT balance of Rs 2,593.05 at 23-Mar-26 (Dr 12,768.48 at 31-Mar-25).',
  'A credit bank balance cannot be shown as cash. Presented as "Book overdraft - Kotak Mahindra Bank" under Other current liabilities (Note 2.5). Confirm against bank statement / BRS whether this is an actual overdraft or unrecorded receipts / timing.'),
 ('Sundry creditors - debit balances', 'Five Sundry Creditor ledgers have DEBIT balances totalling Rs 1,73,500 (Mangala 15,000; Sriguru P V 20,000; Subodh Kumar Mishra 1,21,000; Swami Gangaram 2,500; Yuvaraj 15,000).',
  'Presented as advances to suppliers/consultants under Short-term loans and advances (Note 2.7), not netted against payables. Confirm they are genuine advances and not unrecorded bills or payments posted to wrong ledgers.'),
 ('Unsecured loans', 'Loans (Liability): Samartha Raghava Nagabhushanam Rs 24,80,680 (PY 7,80,680) and Manasa Rs 1,70,180 (PY Nil). No repayment terms, interest or tenure in the TB.',
  'Classified as Long-term borrowings (Note 2.3), following the reference mapping of Directors Loan -> Long-term borrowings. Confirm tenure (if repayable within 12 months they belong under Short-term borrowings), interest (none charged in TB) and the relationship of Manasa (related party?).'),
 ('Related parties', 'Partner Samartha Raghava Nagabhushanam: capital Rs 90,000; loan closing Rs 24,80,680; net loan received during FY 25-26 Rs 17,00,000 (closing less PY closing). Many creditors are individuals grouped as "Consultants".',
  'Only the partner is disclosed as related party (Note 2.11). Gross loan receipts/repayments are not determinable from a TB (only net movement). Confirm whether any consultant/creditor (e.g. Madhavi K Rs 90,000, Sathyanarayana Bhat Rs 64,590) or lender Manasa is a relative of a partner.'),
 ('Trade payables - MSME', 'Sundry Creditors (credit balances) Rs 5,48,320.48 (PY 2,28,000). MSME status of suppliers is not in the TB.',
  'Entire amount shown under "other than micro and small enterprises" pending MSME confirmation (Note 2.4). Schedule III ageing schedule is not prepared: due-date information is not in a TB (and Schedule III does not apply to an LLP).'),
 ('GST', 'GST Payable Rs 7,505.41 (credit) and Input CGST Rs 2,820.94 + Input SGST Rs 2,820.94 (debit) exist as separate ledgers.',
  'Presented gross: GST payable under statutory dues (Note 2.5) and input credits under Other current assets (Note 2.8). Confirm with GSTR-3B / electronic credit ledger whether they should be set off and shown net.'),
 ('Expense classification', 'All expenses are grouped by Tally under "Indirect Expenses". Pooja Expenses Rs 9,59,127.50 and Professional Charges Rs 10,08,252 together exceed revenue of Rs 10,59,363.54. Staff Welfare Rs 7,681 / Internship Charges Rs 12,000 are employee-related; Bank Charges Rs 880.63 could be finance cost.',
  'All expenses are presented under a single "Other expenses" note (2.10) exactly as grouped in Tally, as in the reference. Confirm whether Pooja Expenses are direct cost of services and whether employee-benefit / finance-cost lines should be shown separately.'),
 ('Pre-incorporation expenses', 'Pre-incorporation Expenses Rs 36,000 in FY 25-26 in addition to Rs 4,79,320 in FY 24-25.',
  'Charged to P&L as per TB. Confirm the nature of the FY 25-26 amount (second year after incorporation) and whether any amount should be treated differently.'),
 ('Statutory interest / fees', 'Interest on TDS Rs 2,164 (PY 780) and GST Late Filing Fees Rs 20 are charged to P&L.',
  'Presented as expenses per TB. These are typically disallowed for income-tax; no tax computation has been prepared (the reference tax working sheets were hidden, company-specific and broken with reference errors).'),
 ('Provisions not in TB', 'No audit fee provision, no depreciation/fixed assets, no income-tax provision, no interest on partner loan/capital, no partner remuneration appear in the TB.',
  'Nothing has been added. Current tax is Nil (loss year). If an audit fee or any year-end provision is to be booked, it must first be passed in Tally and the TB re-exported.'),
 ('Going concern / net worth', "Partners' funds are negative: Rs (30,31,142.86) at 23-Mar-26 (PY Rs (10,30,911.52)). Loss for the year Rs 20,00,231.34 on revenue Rs 10,59,363.54.",
  'Reported as per TB. Partners should confirm continued financial support (loans of Rs 26.5 lakh) for a going-concern statement in the accounting policies (Note 1), which is not part of the reference workbook and has not been prepared.'),
 ('Rounding', "Reference presents amounts in Rs '00 (Master!B31 = 100) with display rounding, without ROUND() in formulas (same as reference).",
  "Because underlying values carry paise, in Rs '00 the displayed components of TOTAL LIABILITIES (-30,311 + 26,509 + 5,802 = 2,000) differ by 1 from the displayed total (1,999); the underlying figures balance exactly. Set Master!B31 to 1 to present in full rupees, which removes this display difference."),
 ('Cash', 'Cash in hand Rs 20,754.20 at 23-Mar-26 (PY Nil).', 'Reported as per TB. Confirm physical cash count / cash book at year end.'),
 ('Ledger names', 'Several Tally ledger names contain spelling errors (e.g. "Telepohone Exenses", "Campaining", "Subscrptions", "Accomodation", "Captial").',
  'Exact Tally names are retained in the TB sheet (column A and grouping keys). Only the presentation captions in Note 2.10 have corrected spelling; no amounts are affected.'),
 ('Previous year opening balances', 'The FY 24-25 TB has no brought-forward Profit & Loss balance and includes Pre-incorporation Expenses, indicating FY 24-25 was the first year.',
  "PY opening capital (Note 2.1B) and PY opening reserves (Note 2.2) are therefore Nil and PY 'contributed / received during the year' equals the PY closing balance. Confirm the LLP's date of incorporation."),
 ('Reference sheets not carried', 'Reference workbook contains hidden company income-tax working sheets (IT Depn, IT Comp, DTA, MAT, ARI) and a hidden trade-payables ageing sheet, all with reference errors and specific to a company.',
  'Not reproduced (MAT, EPS, Schedule III ratios and ageing do not apply to an LLP; deferred tax not computed). Master and TB sheets are kept visible here so the mapping can be reviewed; hide them for printing if desired.'),
]
r = 5
for i, (area, obs, req) in enumerate(flags, 1):
    put(fl, f'A{r}', i, h='center', v='top', top=THIN, bottom=THIN, left=THIN, right=THIN)
    put(fl, f'B{r}', area, bold=True, v='top', wrap=True, top=THIN, bottom=THIN, left=THIN, right=THIN)
    put(fl, f'C{r}', obs, v='top', wrap=True, top=THIN, bottom=THIN, left=THIN, right=THIN)
    put(fl, f'D{r}', req, v='top', wrap=True, top=THIN, bottom=THIN, left=THIN, right=THIN)
    fl.row_dimensions[r].height = max(45, 15 * (max(len(obs), len(req)) // 60 + 1))
    r += 1
setup_print(fl, f'Flags!$A$1:$D${r-1}', landscape=True)

wb.active = wb.sheetnames.index('BS')
wb.save(OUT)
print('saved', OUT, 'TB rows', last_tb, 'expense rows', 72, '-', exp_last)
